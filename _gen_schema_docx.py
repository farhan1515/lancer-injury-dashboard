"""Generate LancersPulse_Schema.docx using python-docx."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"C:\STUDIES\sem3\project\Lancer Athlete Hub\v1\LancersPulse\LancersPulse_Schema.docx"

# Schema data: list of (app_name, [(model_name, description, [(field, type, notes), ...])])
SCHEMA = [
    ("accounts", [
        ("Team", "Sports team grouping for players, coaches, and events.", [
            ("id", "AutoField (PK)", "Primary key"),
            ("name", "CharField(100)", "Team name"),
            ("gender", "CharField(1)", "Choices: M=Men, W=Women"),
        ]),
        ("CustomUser", "Extends Django AbstractUser. Inherits username, password, first_name, last_name, email, is_staff, is_active, is_superuser, date_joined, last_login, groups, user_permissions.", [
            ("role", "CharField(10)", "Choices: ADMIN, COACH, DOCTOR, PLAYER. Default PLAYER"),
            ("team", "FK -> Team", "Nullable, on_delete=SET_NULL"),
            ("is_registration_complete", "BooleanField", "Default False"),
            ("phone", "CharField(20)", "Blank"),
            ("gender", "CharField(1)", "Choices: M, F, O, P"),
            ("date_of_birth", "DateField", "Nullable"),
            ("address", "TextField", "Blank"),
            ("city", "CharField(100)", "Blank"),
            ("state", "CharField(100)", "Blank"),
            ("zip_code", "CharField(20)", "Blank"),
            ("country", "CharField(100)", "Default 'USA'"),
            ("emergency_contact_name", "CharField(200)", "Blank"),
            ("emergency_contact_phone", "CharField(20)", "Blank"),
            ("emergency_contact_relationship", "CharField(100)", "Blank"),
            ("emergency_contact_email", "EmailField", "Blank"),
            ("blood_type", "CharField(10)", "Blank"),
            ("medical_conditions", "TextField", "Blank"),
            ("medications", "TextField", "Blank"),
            ("allergies", "TextField", "Blank"),
            ("bio", "TextField", "Blank"),
            ("profile_picture", "ImageField", "upload_to='profile_pictures/'"),
        ]),
        ("PlayerProfile", "Extended athlete profile, one-to-one with CustomUser.", [
            ("user", "OneToOne -> CustomUser", "Cascade"),
            ("preferred_name", "CharField(100)", "Blank"),
            ("alternate_email", "EmailField", "Blank"),
            ("student_id", "CharField(20)", "Blank"),
            ("citizenship", "CharField(100)", "Blank"),
            ("photo_id", "ImageField", "upload_to='player_ids/'"),
            ("local_street_address", "CharField(200)", "Local residence"),
            ("local_city", "CharField(100)", "Local residence"),
            ("local_province", "CharField(100)", "Local residence"),
            ("local_postal_code", "CharField(20)", "Local residence"),
            ("permanent_street_address", "CharField(200)", "Permanent address"),
            ("permanent_city", "CharField(100)", "Permanent address"),
            ("permanent_province_state", "CharField(100)", "Permanent address"),
            ("permanent_postal_zip_code", "CharField(20)", "Permanent address"),
            ("permanent_country", "CharField(100)", "Permanent address"),
            ("sport_gender_category", "CharField(100)", "Blank"),
            ("height_feet", "PositiveIntegerField", "Nullable"),
            ("height_inches", "PositiveIntegerField", "Nullable"),
            ("weight_lbs", "PositiveIntegerField", "Nullable"),
            ("hometown", "CharField(100)", "Blank"),
            ("high_school", "CharField(200)", "Blank"),
            ("high_school_coach_name", "CharField(200)", "Blank"),
            ("position_event", "CharField(100)", "Blank"),
            ("previous_club_team", "CharField(200)", "Blank"),
            ("club_team_coach_name", "CharField(200)", "Blank"),
            ("family_member_oua_usports", "BooleanField", "Default False"),
            ("family_member_details", "TextField", "Blank"),
            ("school_type_last_year", "CharField(100)", "Blank"),
            ("completed_18_credits_last_year", "CharField(10)", "Choices: YES/NO/NA"),
            ("academic_restrictions", "CharField(10)", "Choices: YES/NO"),
            ("faculty", "CharField(200)", "Blank"),
            ("program_of_study", "CharField(200)", "Blank"),
            ("student_status", "CharField(20)", "Choices: UNDERGRADUATE/GRADUATE/OTHER"),
            ("year_of_study", "PositiveIntegerField", "Nullable"),
            ("registered_9_credits_per_term", "CharField(10)", "Choices: YES/NO"),
            ("graduating_this_year", "CharField(10)", "Choices: YES/NO"),
            ("number", "PositiveIntegerField", "Legacy"),
            ("position", "CharField(50)", "Legacy"),
            ("dob", "DateField", "Legacy"),
        ]),
        ("CoachProfile", "Extended coach profile, one-to-one with CustomUser.", [
            ("user", "OneToOne -> CustomUser", "Cascade"),
            ("coaching_experience", "PositiveIntegerField", "Years; nullable"),
            ("specialization", "CharField(100)", "Blank"),
            ("certification", "CharField(100)", "Blank"),
        ]),
        ("DoctorProfile", "Extended doctor profile, one-to-one with CustomUser.", [
            ("user", "OneToOne -> CustomUser", "Cascade"),
            ("medical_license", "CharField(100)", "Blank"),
            ("specialization", "CharField(100)", "Blank"),
            ("years_experience", "PositiveIntegerField", "Nullable"),
        ]),
        ("TeamPermission", "Grants a user (coach/doctor/admin) permissions for an additional team.", [
            ("user", "FK -> CustomUser", "related_name='team_permissions'"),
            ("team", "FK -> Team", "related_name='user_permissions'"),
            ("role_scope", "CharField(10)", "Choices: COACH/DOCTOR/ADMIN"),
            ("created_at", "DateTimeField", "auto_now_add"),
            ("Meta.unique_together", "(user, team, role_scope)", "Composite unique"),
        ]),
        ("TeamPermissionRequest", "A request by a user to gain access to an additional team.", [
            ("user", "FK -> CustomUser", "related_name='team_permission_requests'"),
            ("team", "FK -> Team", "Cascade"),
            ("role_scope", "CharField(10)", "Choices: COACH/DOCTOR/ADMIN"),
            ("justification", "TextField", "Blank"),
            ("status", "CharField(10)", "PENDING/APPROVED/DENIED; default PENDING"),
            ("created_at", "DateTimeField", "auto_now_add"),
            ("reviewed_by", "FK -> CustomUser", "Nullable, SET_NULL"),
            ("reviewed_at", "DateTimeField", "Nullable"),
            ("admin_note", "TextField", "Blank"),
            ("Meta.ordering", "['-created_at']", "Newest first"),
        ]),
        ("EmailRoleMapping", "Maps email domains or specific emails to default roles.", [
            ("email_pattern", "CharField(255)", "Domain or specific email"),
            ("role", "CharField(10)", "Reuses CustomUser.ROLE_CHOICES"),
            ("team", "FK -> Team", "Nullable, SET_NULL"),
            ("is_active", "BooleanField", "Default True"),
        ]),
    ]),
    ("injury_tracking", [
        ("Event", "Team events such as trainings, sessions, and games.", [
            ("team", "FK -> accounts.Team", "related_name='events', cascade"),
            ("created_by", "FK -> CustomUser", "related_name='created_events'"),
            ("event_type", "CharField(20)", "Choices: TRAINING/SESSION/GAME"),
            ("title", "CharField(200)", "Required"),
            ("description", "TextField", "Blank"),
            ("location", "CharField(200)", "Blank"),
            ("start_datetime", "DateTimeField", "Required"),
            ("end_datetime", "DateTimeField", "Required"),
            ("created_at", "DateTimeField", "auto_now_add"),
            ("updated_at", "DateTimeField", "auto_now"),
            ("Meta.ordering", "['start_datetime']", "Earliest first"),
        ]),
        ("InjuryType", "Catalogued types of injuries (e.g. Sprain, Fracture).", [
            ("name", "CharField(100)", "Unique"),
            ("description", "TextField", "Blank"),
        ]),
        ("BodyPart", "Catalogued body parts that can be injured.", [
            ("name", "CharField(100)", "Unique"),
        ]),
        ("InjurySeverity", "Severity levels with display colour.", [
            ("name", "CharField(50)", "Unique"),
            ("color_code", "CharField(7)", "Hex (e.g. #FF0000)"),
            ("description", "TextField", "Blank"),
        ]),
        ("InjuryRecord", "Core injury record. Drives the analytics dashboard.", [
            ("player", "FK -> CustomUser", "related_name='injuries', cascade"),
            ("reported_by", "FK -> CustomUser", "related_name='reported_injuries'"),
            ("injury_date", "DateField", "Required"),
            ("reported_date", "DateTimeField", "auto_now_add"),
            ("injury_type", "FK -> InjuryType", "Cascade"),
            ("body_part", "FK -> BodyPart", "Cascade"),
            ("severity", "FK -> InjurySeverity", "Cascade"),
            ("status", "CharField(20)", "ACTIVE/RECOVERING/RECOVERED/CHRONIC; default ACTIVE"),
            ("description", "TextField", "Required"),
            ("symptoms", "TextField", "Blank"),
            ("treatment", "CharField(20)", "REST/PHYSIO/SURGERY/MEDICATION/OTHER"),
            ("treatment_notes", "TextField", "Blank"),
            ("contact_type", "CharField(20)", "CONTACT/NON_CONTACT/OVERUSE/UNKNOWN; default UNKNOWN"),
            ("missed_games", "PositiveIntegerField", "Nullable"),
            ("missed_practices", "PositiveIntegerField", "Nullable"),
            ("season_year", "PositiveIntegerField", "Nullable; auto-populated from injury_date (Sept cutoff)"),
            ("estimated_recovery_time", "PositiveIntegerField", "Days; nullable"),
            ("actual_recovery_time", "PositiveIntegerField", "Days; nullable"),
            ("return_to_play_date", "DateField", "Nullable"),
            ("requires_surgery", "BooleanField", "Default False"),
            ("surgery_date", "DateField", "Nullable"),
            ("medical_clearance", "BooleanField", "Default False"),
            ("clearance_date", "DateField", "Nullable"),
            ("follow_up_required", "BooleanField", "Default False"),
            ("follow_up_date", "DateField", "Nullable"),
            ("follow_up_notes", "TextField", "Blank"),
            ("is_confidential", "BooleanField", "Default False"),
            ("created_at", "DateTimeField", "auto_now_add"),
            ("updated_at", "DateTimeField", "auto_now"),
            ("Meta.ordering", "['-injury_date']", "Newest first"),
            ("Meta.permissions", "view_own_injuries, view_team_injuries, view_all_injuries", "Custom perms"),
            ("(property) days_since_injury", "int", "today - injury_date"),
            ("(property) is_fully_recovered", "bool", "status==RECOVERED and medical_clearance"),
            ("(property) is_preventable", "bool", "contact_type in (NON_CONTACT, OVERUSE)"),
            ("(property) total_missed_events", "int|None", "missed_games + missed_practices"),
        ]),
        ("InjuryFollowUp", "Follow-up records linked to an InjuryRecord.", [
            ("injury", "FK -> InjuryRecord", "related_name='follow_ups', cascade"),
            ("follow_up_date", "DateField", "Required"),
            ("notes", "TextField", "Required"),
            ("status_update", "CharField(20)", "Reuses InjuryRecord.STATUS_CHOICES"),
            ("created_by", "FK -> CustomUser", "Cascade"),
            ("created_at", "DateTimeField", "auto_now_add"),
        ]),
        ("TeamRoster", "Roster membership linking players to teams.", [
            ("team", "FK -> accounts.Team", "related_name='roster'"),
            ("player", "FK -> CustomUser", "related_name='team_memberships'"),
            ("position", "CharField(50)", "Blank"),
            ("jersey_number", "PositiveIntegerField", "Nullable"),
            ("is_active", "BooleanField", "Default True"),
            ("joined_date", "DateField", "auto_now_add"),
            ("Meta.unique_together", "(team, player)", "Composite unique"),
        ]),
        ("InjuryAnalytics", "Per-team, per-season precomputed aggregates.", [
            ("team", "FK -> accounts.Team", "Cascade"),
            ("season_year", "PositiveIntegerField", "Required"),
            ("total_injuries", "PositiveIntegerField", "Default 0"),
            ("active_injuries", "PositiveIntegerField", "Default 0"),
            ("recovered_injuries", "PositiveIntegerField", "Default 0"),
            ("most_common_injury_type", "FK -> InjuryType", "Nullable, SET_NULL"),
            ("most_common_body_part", "FK -> BodyPart", "Nullable, SET_NULL"),
            ("average_recovery_time", "FloatField", "Nullable"),
            ("created_at", "DateTimeField", "auto_now_add"),
            ("Meta.unique_together", "(team, season_year)", "Composite unique"),
        ]),
    ]),
    ("injuries", [
        ("InjuryReport", "Alternate injury submission model (parallel to InjuryRecord).", [
            ("player", "FK -> AUTH_USER_MODEL", "related_name='injury_reports', cascade"),
            ("doctor", "FK -> AUTH_USER_MODEL", "related_name='submitted_reports', SET_NULL"),
            ("reported_date", "DateField", "auto_now_add"),
            ("injury_date", "DateField", "null=True, blank=False"),
            ("body_part", "CharField(20)", "HEAD/NECK/SHOULDER/ARM/HAND/CHEST/BACK/HIP/THIGH/KNEE/LEG/ANKLE/FOOT/OTHER"),
            ("diagnosis", "CharField(255)", "Required"),
            ("mechanism", "CharField(255)", "Blank"),
            ("severity", "CharField(10)", "MINOR/MODERATE/SEVERE"),
            ("imaging_required", "BooleanField", "Default False"),
            ("imaging_type", "CharField(10)", "XRAY/MRI/CT/US/NONE"),
            ("imaging_details", "CharField(255)", "Blank"),
            ("treatment_given", "TextField", "Blank"),
            ("recommended_followup", "TextField", "Blank"),
            ("time_lost_days", "IntegerField", "Nullable"),
            ("expected_return_date", "DateField", "Nullable"),
            ("restrictions", "TextField", "Blank"),
            ("notes", "TextField", "Blank"),
            ("Meta.ordering", "['-reported_date']", "Newest first"),
        ]),
    ]),
]

RELATIONSHIP_DIAGRAM = """Team ---< CustomUser.team
     ---< Event.team
     ---< TeamRoster.team
     ---< TeamPermission.team
     ---< TeamPermissionRequest.team
     ---< EmailRoleMapping.team
     ---< InjuryAnalytics.team

CustomUser ---< InjuryRecord.player / .reported_by
           ---< InjuryReport.player / .doctor
           ---< Event.created_by
           ---< InjuryFollowUp.created_by
           ---< TeamRoster.player
           ---< TeamPermission.user
           ---< TeamPermissionRequest.user / .reviewed_by
           ---1:1--- PlayerProfile / CoachProfile / DoctorProfile

InjuryType     ---< InjuryRecord.injury_type, InjuryAnalytics.most_common_injury_type
BodyPart       ---< InjuryRecord.body_part,  InjuryAnalytics.most_common_body_part
InjurySeverity ---< InjuryRecord.severity
InjuryRecord   ---< InjuryFollowUp.injury
"""


def shade_cell(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tc_pr.append(shd)


def add_model_table(doc, fields):
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Light Grid Accent 1'
    table.autofit = False
    widths = [Inches(2.2), Inches(1.9), Inches(2.7)]
    hdr = table.rows[0].cells
    headers = ["Field", "Type", "Constraints / Notes"]
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(11)
        hdr[i].width = widths[i]
        shade_cell(hdr[i], "1F3864")
        for r in hdr[i].paragraphs[0].runs:
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for field, ftype, notes in fields:
        row = table.add_row().cells
        row[0].text = field
        row[1].text = ftype
        row[2].text = notes
        for i in range(3):
            row[i].width = widths[i]
            for p in row[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)


def main():
    doc = Document()

    # Default style
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    # Title
    title = doc.add_heading('LancersPulse — Database Schema', level=0)
    sub = doc.add_paragraph()
    sub_run = sub.add_run('Django models across the accounts, injury_tracking, and injuries apps.')
    sub_run.italic = True
    sub_run.font.size = Pt(11)

    doc.add_paragraph(
        "This document lists every model in the LancersPulse project with its fields, "
        "data types, and constraints. Foreign-key relationships and Meta options are noted "
        "in the Constraints / Notes column. A relationship summary diagram follows the model tables."
    )

    for app_name, models in SCHEMA:
        doc.add_heading(f"App: {app_name}", level=1)
        for model_name, description, fields in models:
            doc.add_heading(model_name, level=2)
            if description:
                p = doc.add_paragraph()
                run = p.add_run(description)
                run.italic = True
            add_model_table(doc, fields)
            doc.add_paragraph()

    doc.add_heading("Relationship Summary", level=1)
    doc.add_paragraph(
        "The diagram below shows foreign-key relationships across the project. "
        "'---<' denotes a one-to-many relationship; '---1:1---' denotes a one-to-one relationship."
    )
    code_p = doc.add_paragraph()
    code_run = code_p.add_run(RELATIONSHIP_DIAGRAM)
    code_run.font.name = 'Consolas'
    code_run.font.size = Pt(9)

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
